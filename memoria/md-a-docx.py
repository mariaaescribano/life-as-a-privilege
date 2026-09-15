#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Convierte MEMORIA-COMPLETA.md en un .docx que hereda los estilos de la
plantilla oficial de la EPS.

USO:
    python md-a-docx.py

    python md-a-docx.py --plantilla "C:\\ruta\\plantilla.docx" --salida MEMORIA.docx

Qué hace:
  · Usa la plantilla de la EPS como base: conserva sus estilos (Normal,
    Heading 1/2, Caption, «Fuente figua-tabla», Table Grid), sus márgenes y
    su tamaño de página, pero borra todo su contenido.
  · Títulos de nivel 1 → «Título 1» con salto de página delante.
  · Tablas de Markdown → tablas de Word con estilo Table Grid y cabecera en negrita.
  · Pies de figura y títulos de tabla → estilo «Caption»; la línea de fuente,
    estilo «Fuente figua-tabla».
  · Las imágenes ![...](figuras/...) se insertan de verdad, escaladas al ancho.
  · Los huecos [FIGURA n: ...] se convierten en un recuadro visible que hay
    que sustituir por la imagen.
  · Los comentarios <!-- ... --> se convierten en notas en gris, para que se
    vean en Word y puedas borrarlas al terminar.

Requisito: pip install python-docx
"""

import argparse
import os
import re
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
PLANTILLA_POR_DEFECTO = os.path.join(
    os.path.expanduser("~"), "Downloads",
    "2022 Guia para el desarrollo de TFG-multimedia.docx",
)
ANCHO_MAX = Cm(15.0)          # ancho útil de la página con los márgenes de la EPS
GRIS = RGBColor(0x80, 0x80, 0x80)


# ─────────────────────────── utilidades de estilo ───────────────────────────

def vaciar_cuerpo(doc):
    """Borra todo el contenido de la plantilla conservando estilos y sección.

    Además suelta las imágenes de la plantilla: si no se quitan sus relaciones,
    las 31 figuras del documento original viajan dentro del .docx aunque no se
    vean, y engordan el fichero varios megas.
    """
    body = doc.element.body
    for hijo in list(body):
        if hijo.tag.endswith("}sectPr"):   # propiedades de página: se conservan
            continue
        body.remove(hijo)

    parte = doc.part
    for rId, rel in list(parte.rels.items()):
        if "image" in rel.reltype:
            parte.drop_rel(rId)


def asegurar_estilo(doc, nombre, basado_en=None, tam=None, negrita=None):
    """Crea un estilo de párrafo si la plantilla no lo trae."""
    try:
        return doc.styles[nombre]
    except KeyError:
        pass
    estilo = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
    if basado_en:
        try:
            estilo.base_style = doc.styles[basado_en]
        except KeyError:
            pass
    if tam is not None:
        estilo.font.size = Pt(tam)
    if negrita is not None:
        estilo.font.bold = negrita
    return estilo


def salto_de_pagina_antes(parrafo):
    parrafo.paragraph_format.page_break_before = True


def ocupar_ancho_de_pagina(tabla):
    """La tabla ocupa el ancho útil y reparte columnas según su contenido."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tabla.autofit = True
    tblPr = tabla._tbl.tblPr
    for etiqueta, atributos in (
        ("w:tblW", {"w:type": "pct", "w:w": "5000"}),   # 5000 = 100 %
        ("w:tblLayout", {"w:type": "autofit"}),
    ):
        el = OxmlElement(etiqueta)
        for k, v in atributos.items():
            el.set(qn(k), v)
        tblPr.append(el)


def repetir_cabecera(fila):
    """Marca la fila como cabecera: si la tabla parte de página, se repite."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    trPr = fila._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


# ─────────────────────────── formato en línea ───────────────────────────

# **negrita**, *cursiva*, `código`, [texto](enlace)
PATRON = re.compile(
    r"(\*\*.+?\*\*)"      # negrita
    r"|(\*[^*\n]+?\*)"    # cursiva
    r"|(`[^`\n]+?`)"      # código
    r"|(\[[^\]]+?\]\([^)]+?\))"   # enlace
)


def escribir_texto(parrafo, texto):
    """Escribe texto aplicando negrita, cursiva, código y enlaces."""
    pos = 0
    for m in PATRON.finditer(texto):
        if m.start() > pos:
            parrafo.add_run(texto[pos:m.start()])
        trozo = m.group(0)
        if trozo.startswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`"):
            r = parrafo.add_run(trozo[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif trozo.startswith("["):
            etiqueta, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", trozo).groups()
            if url.startswith(("http://", "https://")):
                r = parrafo.add_run(f"{etiqueta} ({url})")
            else:
                r = parrafo.add_run(etiqueta)
            r.underline = True
        else:  # cursiva
            parrafo.add_run(trozo[1:-1]).italic = True
        pos = m.end()
    if pos < len(texto):
        parrafo.add_run(texto[pos:])


# ─────────────────────────── conversión ───────────────────────────

def convertir(ruta_md, ruta_plantilla, ruta_salida):
    with open(ruta_md, encoding="utf-8") as f:
        lineas = f.read().split("\n")

    if os.path.exists(ruta_plantilla):
        doc = Document(ruta_plantilla)
        vaciar_cuerpo(doc)
        print(f"Plantilla: {ruta_plantilla}")
    else:
        doc = Document()
        print("AVISO: no se encontró la plantilla de la EPS; se usa la de Word.")
        print("       Los estilos habrá que ajustarlos a mano.")

    asegurar_estilo(doc, "Heading 3", basado_en="Heading 2", tam=12, negrita=True)
    estilo_pie = "Fuente figua-tabla" if "Fuente figua-tabla" in [s.name for s in doc.styles] else "Caption"

    i = 0
    n_tablas = n_figuras = n_huecos = n_notas = 0
    primer_titulo = True

    while i < len(lineas):
        linea = lineas[i]
        pelada = linea.strip()

        # ── comentarios HTML: se vuelcan como nota gris ──
        if pelada.startswith("<!--"):
            bloque = []
            while i < len(lineas) and "-->" not in lineas[i]:
                bloque.append(lineas[i]); i += 1
            if i < len(lineas):
                bloque.append(lineas[i]); i += 1
            texto = " ".join(bloque).replace("<!--", "").replace("-->", "")
            texto = re.sub(r"\s+", " ", texto).strip()
            texto = re.sub(r"^[─\s]+", "", texto)
            if texto:
                p = doc.add_paragraph()
                r = p.add_run("NOTA (borrar antes de entregar): " + texto)
                r.italic = True
                r.font.color.rgb = GRIS
                r.font.size = Pt(9)
                n_notas += 1
            continue

        # ── línea en blanco ──
        if not pelada:
            i += 1
            continue

        # ── separador horizontal: se ignora ──
        if pelada in ("---", "***", "___"):
            i += 1
            continue

        # ── títulos ──
        m = re.match(r"^(#{1,6})\s+(.*)$", pelada)
        if m:
            nivel, texto = len(m.group(1)), m.group(2).strip()
            estilo = {1: "Heading 1", 2: "Heading 2"}.get(nivel, "Heading 3")
            p = doc.add_paragraph(texto, style=estilo)
            if nivel == 1 and not primer_titulo:
                salto_de_pagina_antes(p)
            if nivel == 1:
                primer_titulo = False
            i += 1
            continue

        # ── imagen ──
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", pelada)
        if m:
            ruta_img = os.path.join(AQUI, m.group(2))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(ruta_img):
                p.add_run().add_picture(ruta_img, width=ANCHO_MAX)
                n_figuras += 1
            else:
                r = p.add_run(f"[FALTA LA IMAGEN: {m.group(2)}]")
                r.bold = True
                r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            i += 1
            continue

        # ── hueco de figura o de tabla ──
        m = re.match(r"^\[(FIGURA|TABLA)\s+(\d+)\s*:\s*(.*)\]\s*$", pelada)
        if m:
            tipo, num, desc = m.group(1), m.group(2), m.group(3)
            if tipo == "TABLA":
                i += 1          # el título de la tabla va justo debajo; no hace falta marca
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"[ SUSTITUIR POR LA FIGURA {num}: {desc} ]")
            r.bold = True
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            n_huecos += 1
            i += 1
            continue

        # ── pie de figura / título de tabla: línea entera en cursiva ──
        if re.match(r"^\*(Figura|Tabla)\s+\d+\..*\*$", pelada):
            doc.add_paragraph(pelada.strip("*"), style="Caption")
            i += 1
            continue

        if re.match(r"^\*\(Fuente.*\)\*$", pelada):
            doc.add_paragraph(pelada.strip("*"), style=estilo_pie)
            i += 1
            continue

        # ── bloque de código ──
        if pelada.startswith("```"):
            i += 1
            codigo = []
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                codigo.append(lineas[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            r = p.add_run("\n".join(codigo))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue

        # ── tabla ──
        if pelada.startswith("|") and i + 1 < len(lineas) and re.match(r"^\|[\s:\-|]+\|$", lineas[i + 1].strip()):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                fila = lineas[i].strip()
                if not re.match(r"^\|[\s:\-|]+\|$", fila):
                    filas.append([c.strip() for c in fila.strip("|").split("|")])
                i += 1
            if filas:
                n_cols = max(len(f) for f in filas)
                tabla = doc.add_table(rows=0, cols=n_cols)
                try:
                    tabla.style = doc.styles["Table Grid"]
                except KeyError:
                    pass
                ocupar_ancho_de_pagina(tabla)
                for idx, fila in enumerate(filas):
                    wfila = tabla.add_row()
                    for j in range(n_cols):
                        texto = fila[j] if j < len(fila) else ""
                        par = wfila.cells[j].paragraphs[0]
                        # a la izquierda, no justificado: en columnas estrechas
                        # el justificado abre huecos enormes entre palabras
                        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        escribir_texto(par, texto)
                        for run in par.runs:
                            run.font.size = Pt(9)
                            if idx == 0:
                                run.bold = True
                    if idx == 0:
                        repetir_cabecera(wfila)
                doc.add_paragraph()
                n_tablas += 1
            continue

        # ── cita ──
        if pelada.startswith(">"):
            texto = pelada.lstrip("> ").strip()
            if texto:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.5)
                escribir_texto(p, texto)
                for r in p.runs:
                    r.italic = True
            i += 1
            continue

        # ── viñetas y listas numeradas ──
        m = re.match(r"^[-*]\s+(.*)$", pelada)
        if m:
            p = doc.add_paragraph(style="List Bullet" if "List Bullet" in [s.name for s in doc.styles] else None)
            if p.style.name == "Normal":
                p.paragraph_format.left_indent = Cm(0.8)
                p.add_run("· ")
            escribir_texto(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", pelada)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.add_run(f"{m.group(1)}. ")
            escribir_texto(p, m.group(2))
            i += 1
            continue

        # ── párrafo normal ──
        p = doc.add_paragraph()
        escribir_texto(p, pelada)
        i += 1

    doc.save(ruta_salida)
    print()
    print(f"Generado: {ruta_salida}")
    print(f"  tablas insertadas ........ {n_tablas}")
    print(f"  imágenes insertadas ...... {n_figuras}")
    print(f"  huecos de figura ......... {n_huecos}  (en rojo, hay que sustituirlos)")
    print(f"  notas para la autoría .... {n_notas}  (en gris, hay que borrarlas)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", default=os.path.join(AQUI, "MEMORIA-COMPLETA.md"))
    ap.add_argument("--plantilla", default=PLANTILLA_POR_DEFECTO)
    ap.add_argument("--salida", default=os.path.join(AQUI, "MEMORIA.docx"))
    args = ap.parse_args()

    if not os.path.exists(args.md):
        sys.exit(f"No se encuentra el Markdown: {args.md}")
    convertir(args.md, args.plantilla, args.salida)
